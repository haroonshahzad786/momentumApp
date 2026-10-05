"""Regenerate the Traceability .docx from PROGRESS_AND_TRACEABILITY.md.

    python md_to_docx.py   (run from this folder)

Handles the subset of Markdown the doc uses: headings, bullets (2 levels),
numbered items, block quotes, tables, images, rules, **bold**, *italic*, `code`.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).parent
SRC = HERE / "PROGRESS_AND_TRACEABILITY.md"
OUT = HERE / "Moore Momentum - Build Progress & Traceability.docx"

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
        elif part.startswith("[") and "](" in part:
            par.add_run(part[1:part.index("](")])
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    width = max(len(r) for r in cells)
    t = doc.add_table(rows=len(cells), cols=width)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        for j in range(width):
            p = t.cell(i, j).paragraphs[0]
            add_runs(p, r[j] if j < len(r) else "")
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            table(doc, rows)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        img = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", s)
        if m:
            add_runs(doc.add_heading(level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0), m.group(2))
        elif img:
            path = HERE / img.group(2)
            if path.exists():
                doc.add_picture(str(path), width=Inches(5.5))
                cap = doc.add_paragraph()
                r = cap.add_run(img.group(1))
                r.italic = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        elif s == "---":
            pass  # section breaks are carried by the headings
        elif s.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            add_runs(p, s[2:])
        elif re.match(r"^- ", line):
            add_runs(doc.add_paragraph(style="List Bullet"), s[2:])
        elif re.match(r"^\s{2,}- ", line):
            add_runs(doc.add_paragraph(style="List Bullet 2"), s[2:])
        elif re.match(r"^\d+\. ", s):
            add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\.\s", "", s))
        else:
            add_runs(doc.add_paragraph(), s)
        i += 1
    doc.save(OUT)
    print(f"wrote {OUT.name}")


if __name__ == "__main__":
    main()
